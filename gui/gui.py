# gui/gui.py - Главное окно приложения (CustomTkinter, тёмная тема)
#
# Структура:
#   - Сайдбар слева: разделы «Штрихкоды» / «QR-паспорта» + инструменты
#   - Контент: страница раздела (крупный заголовок, главная кнопка, рабочая область)
#   - Нижняя панель: статус и счётчики
# Все цвета, шрифты и виджеты — из gui/theme.py.

import tkinter as tk
from tkinter import filedialog, messagebox
from tkinter import StringVar, BooleanVar, IntVar
import customtkinter as ctk
import sys
import os
import re
import random
import uuid
import queue
import tempfile
import threading
import traceback
from datetime import datetime

from . import theme as T
from .gui_widgets import ContextMenu
from .gui_preview import PreviewWindow
from .gui_template_editor import TemplateEditor
from .gui_qr_tab import QRPassportTab

from settings import SettingsManager
from barcode_generator import BarcodeGenerator
from utils import BARCODE_TYPES, open_pdf, validate_code, read_codes_with_metadata

from logger import logger

T.apply_theme()


# ========== ДЕКОРАТОР ДЛЯ ОБРАБОТКИ ОШИБОК ==========
def error_handler(func):
    """Декоратор для обработки ошибок в методах GUI"""
    def wrapper(self, *args, **kwargs):
        try:
            return func(self, *args, **kwargs)
        except Exception as e:
            error_msg = traceback.format_exc()
            logger.error(f"Ошибка в {func.__name__}:\n{error_msg}")

            try:
                messagebox.showerror(
                    "Ошибка",
                    f"Произошла ошибка при выполнении операции:\n\n{str(e)}\n\n"
                    f"Ошибка записана в лог-файл: barcode_app.log"
                )
            except Exception:
                pass

            try:
                self.set_status(f"Ошибка: {str(e)[:50]}", 'error')
            except Exception:
                pass

            return None
    return wrapper


def plural(n, one, few, many):
    """1 код, 2 кода, 5 кодов."""
    n = abs(n) % 100
    if 11 <= n <= 19:
        return many
    n %= 10
    if n == 1:
        return one
    if 2 <= n <= 4:
        return few
    return many


# Названия типов без эмодзи: «Code 128 (универсальный)» и короткое «Code 128»
BARCODE_TYPE_LABELS = {val: label.split(" ", 1)[1] for val, label in BARCODE_TYPES}
BARCODE_TYPE_SHORT = {val: label.split(" (")[0] for val, label in BARCODE_TYPE_LABELS.items()}

# Попробуем импортировать tkinterdnd2
try:
    from tkinterdnd2 import DND_FILES
    DND_AVAILABLE = True
    logger.info("✅ Drag & Drop поддержка включена (tkinterdnd2)")
except ImportError:
    DND_AVAILABLE = False
    logger.warning("⚠️ Drag & Drop не поддерживается (tkinterdnd2 не установлен)")


class BarcodeApp:
    """Главный класс приложения"""

    def __init__(self, root):
        self.root = root
        self.root.title("barcodeGen")
        T.fit_to_screen(self.root, 1280, 820, min_width=980, min_height=620)

        if isinstance(self.root, ctk.CTk):
            self.root.configure(fg_color=T.BG)
        else:
            self.root.configure(bg=T.BG)

        logger.info("🚀 Запуск приложения")

        self.settings = SettingsManager()

        self.generator = BarcodeGenerator(self.settings)
        self.generator.set_status_callback(self.set_status)

        self.barcode_type = StringVar(value=self.settings.get("barcode_type", "code128"))
        self.label_width = StringVar(value=self.settings.get("barcode_label_width", "58"))
        self.label_height = StringVar(value=self.settings.get("barcode_label_height", "40"))
        self.margin_top = StringVar(value=self.settings.get("barcode_margin_top", "5"))
        self.margin_bottom = StringVar(value=self.settings.get("barcode_margin_bottom", "5"))
        self.margin_left = StringVar(value=self.settings.get("barcode_margin_left", "5"))
        self.margin_right = StringVar(value=self.settings.get("barcode_margin_right", "5"))
        self.show_border = BooleanVar(value=self.settings.get("barcode_show_border", False))
        self.labels_per_page = IntVar(value=self.settings.get("barcode_labels_per_page", 1))
        self.selected_preset = StringVar(value=self.settings.get("selected_preset", "58×40 мм (стандарт для чеков)"))
        self.selected_only = BooleanVar(value=self.settings.get("barcode_selected_only", False))
        self.batch_mode = BooleanVar(value=self.settings.get("barcode_batch_mode", False))
        self.use_template = BooleanVar(value=self.settings.get("barcode_use_template", True))
        self.show_text = BooleanVar(value=self.settings.get("barcode_show_text", True))
        self.barcode_dpi = StringVar(value=self.settings.get("barcode_dpi", "600"))
        self.search_var = StringVar(value="")

        self.history_data = []
        # Данные из CSV/Excel для плейсхолдеров шаблона: {код: {name, price, ...}}
        self.codes_metadata = dict(self.settings.get("codes_metadata") or {})
        self._job = None  # идущая фоновая генерация
        self.recent_files = []
        self.preview_selection = []

        self.progress_bar = None
        self.progress_window = None
        self.nav_buttons = {}
        self.current_page = "input"
        self.qr_tab = None

        self.create_widgets()
        self.setup_shortcuts()
        self.setup_auto_save()

        if DND_AVAILABLE:
            self.setup_drag_drop()

        self.root.protocol("WM_DELETE_WINDOW", self.on_closing)
        self.root.after(500, self.load_backup)
        self.load_history()

        logger.info("✅ Приложение успешно инициализировано")

    # ========== ОБРАБОТКА ЗАКРЫТИЯ ==========
    def on_closing(self):
        try:
            logger.info("Закрытие приложения пользователем")
            if self._job is not None:
                self._job["cancel"].set()
            self.save_current_settings()

            if self.text_codes.get("1.0", "end").strip():
                if messagebox.askyesno("Сохранение",
                        "У вас есть несохранённые коды. Сохранить перед выходом?"):
                    self.save_to_file()

            self.settings.clear_temp_codes()
            logger.info("✅ Приложение закрыто успешно")
        except Exception as e:
            logger.error(f"Ошибка при закрытии: {e}")
        finally:
            self.root.destroy()

    # ========== ГОРЯЧИЕ КЛАВИШИ ==========
    def setup_shortcuts(self):
        self.root.bind('<Control-o>', lambda e: self.load_from_file())
        self.root.bind('<Control-s>', lambda e: self.save_to_file())
        self.root.bind('<Control-g>', lambda e: self.generate_pdf())
        self.root.bind('<Control-r>', lambda e: self.generate_random_code())
        self.root.bind('<Control-f>', lambda e: self.focus_search())
        self.root.bind('<F5>', lambda e: self.clear_codes())
        self.root.bind('<Control-h>', lambda e: self.show_history_window())
        self.root.bind('<Control-p>', lambda e: self.show_preview_window())
        self.root.bind('<Control-t>', lambda e: self.show_template_dialog())
        logger.info("✅ Горячие клавиши настроены")

    def _bind_paste_shortcut(self, ctk_textbox):
        targets = [ctk_textbox]
        if hasattr(ctk_textbox, "_textbox"):
            targets.append(ctk_textbox._textbox)
        for widget in targets:
            widget.bind("<Control-v>", self._paste_into_codes)
            widget.bind("<Control-V>", self._paste_into_codes)

    def _paste_into_codes(self, event=None):
        try:
            clip = self.root.clipboard_get()
        except Exception:
            return "break"
        try:
            self.text_codes.delete("sel.first", "sel.last")
        except Exception:
            pass
        self.text_codes.insert("insert", clip)
        self.update_all()
        return "break"

    def focus_search(self):
        self.show_page("input")
        if hasattr(self, 'search_entry'):
            self.search_entry.focus_set()

    # ========== АВТОСОХРАНЕНИЕ ==========
    def setup_auto_save(self):
        def auto_save():
            if self.text_codes.get("1.0", "end").strip():
                self.settings.save_temp_codes(self.text_codes.get("1.0", "end"))
                self.set_status("Автосохранение выполнено", 'success')
                self.root.after(2000, lambda: self.update_status_text())
            self.root.after(30000, auto_save)
        self.root.after(30000, auto_save)
        logger.info("✅ Автосохранение настроено (30 секунд)")

    @error_handler
    def load_backup(self):
        try:
            if self.root.winfo_exists():
                temp_codes = self.settings.load_temp_codes()
                if temp_codes:
                    if messagebox.askyesno("Восстановление", "Обнаружены несохранённые коды. Восстановить?"):
                        self.text_codes.delete("1.0", "end")
                        self.text_codes.insert("1.0", temp_codes)
                        self.update_all()
                        self.set_status("Восстановлено из автосохранения", 'success')
                        logger.info("Восстановлены коды из автосохранения")
        except Exception as e:
            logger.error(f"Ошибка восстановления: {e}")

    # ========== DRAG & DROP ==========
    def setup_drag_drop(self):
        target = self.text_codes
        if hasattr(self.text_codes, "_textbox"):
            target = self.text_codes._textbox
        try:
            target.drop_target_register(DND_FILES)
            target.dnd_bind('<<Drop>>', self.on_drop)
            logger.info("✅ Drag & Drop настроен")
        except Exception as e:
            logger.warning(f"⚠️ Не удалось настроить Drag & Drop: {e}")

    @error_handler
    def on_drop(self, event):
        # splitlist корректно разбирает пути с пробелами: {C:/Мои файлы/codes.txt}
        for file_path in self.root.tk.splitlist(event.data):
            if file_path.lower().endswith(('.txt', '.csv', '.xlsx')):
                self.load_codes_from_file(file_path)
                break

    # ========== СОЗДАНИЕ ИНТЕРФЕЙСА ==========
    def create_widgets(self):
        self.root.grid_rowconfigure(0, weight=1)
        self.root.grid_columnconfigure(1, weight=1)

        self.create_sidebar(self.root).grid(row=0, column=0, sticky="nsw")

        self.pages_container = ctk.CTkFrame(self.root, fg_color=T.BG, corner_radius=0)
        self.pages_container.grid(row=0, column=1, sticky="nsew")
        self.pages_container.grid_rowconfigure(0, weight=1)
        self.pages_container.grid_columnconfigure(0, weight=1)

        self.input_page = ctk.CTkFrame(self.pages_container, fg_color=T.BG, corner_radius=0)
        self.qr_page = ctk.CTkFrame(self.pages_container, fg_color=T.BG, corner_radius=0)

        self.create_input_page(self.input_page)

        # Два фрейма в одной ячейке grid — переключение через tkraise()
        self.input_page.grid(row=0, column=0, sticky="nsew")
        self.qr_page.grid(row=0, column=0, sticky="nsew")

        self.create_status_bar(self.root).grid(row=1, column=0, columnspan=2, sticky="ew")

        self.show_page("input")
        self.update_all()

    def _nav_button(self, parent, text, command, icon_name):
        btn = ctk.CTkButton(
            parent, text=f"  {text.upper()}", anchor="w", command=command,
            image=T.icon(icon_name, 20, T.SIDEBAR_TEXT_2), compound="left",
            fg_color="transparent", hover_color=T.SIDEBAR_HOVER, text_color=T.SIDEBAR_TEXT_2,
            font=T.font(13, "extrabold"), corner_radius=12, height=46,
        )
        btn._icon_name = icon_name
        return btn

    def create_sidebar(self, parent):
        sidebar = ctk.CTkFrame(parent, fg_color=T.SIDEBAR, corner_radius=0, width=240)
        sidebar.grid_propagate(False)
        sidebar.pack_propagate(False)

        logo_frame = ctk.CTkFrame(sidebar, fg_color="transparent")
        logo_frame.pack(fill="x", padx=20, pady=(24, 28))

        # Логотип-мозаика: капля из разноцветных бутылок
        logo_tile = ctk.CTkFrame(logo_frame, fg_color="#FFFFFF", corner_radius=14,
                                 width=54, height=54)
        logo_tile.pack(side="left")
        logo_tile.pack_propagate(False)
        ctk.CTkLabel(logo_tile, image=T.logo_image(46), text="").place(
            relx=0.5, rely=0.5, anchor="center")

        names = ctk.CTkFrame(logo_frame, fg_color="transparent")
        names.pack(side="left", padx=(12, 0))
        T.label(names, "BARCODEGEN", size=16, weight="black", color=T.SIDEBAR_TEXT,
                height=20).pack(anchor="w")
        T.label(names, "ЭТИКЕТКИ И QR", size=10, weight="extrabold", color=T.YELLOW,
                height=14).pack(anchor="w")

        nav_items = [
            ("input", "Штрихкоды", "barcode"),
            ("qr", "QR-паспорта", "qr"),
        ]
        for key, text, icon_name in nav_items:
            btn = self._nav_button(sidebar, text, lambda k=key: self.show_page(k), icon_name)
            btn.pack(fill="x", padx=12, pady=2)
            self.nav_buttons[key] = btn

        T.overline(sidebar, "Инструменты", color=T.SIDEBAR_TEXT_2).pack(anchor="w", padx=26, pady=(28, 6))
        self._nav_button(sidebar, "Шаблон этикетки", self.show_template_dialog, "template"
                         ).pack(fill="x", padx=12, pady=2)
        self._nav_button(sidebar, "История", self.show_history_window, "history"
                         ).pack(fill="x", padx=12, pady=2)

        hints = ctk.CTkFrame(sidebar, fg_color="transparent")
        hints.pack(side="bottom", anchor="w", padx=26, pady=(0, 22))
        for row, (keys, action) in enumerate([("Ctrl+G", "создать PDF"),
                                              ("Ctrl+P", "предпросмотр"),
                                              ("Ctrl+O", "открыть файл"),
                                              ("Ctrl+F", "поиск")]):
            T.label(hints, keys, size=11, weight="extrabold", color=T.SIDEBAR_TEXT, height=20
                    ).grid(row=row, column=0, sticky="w", padx=(0, 10))
            T.label(hints, action, size=11, color=T.SIDEBAR_TEXT_2, height=20
                    ).grid(row=row, column=1, sticky="w")

        return sidebar

    def show_page(self, name):
        self.current_page = name

        if name == "input":
            self.input_page.tkraise()
        elif name == "qr":
            if self.qr_tab is None:
                self.qr_tab = QRPassportTab(self.qr_page, self)
            self.qr_page.tkraise()

        for key, btn in self.nav_buttons.items():
            active = key == name
            btn.configure(
                fg_color=T.YELLOW if active else "transparent",
                hover_color=T.YELLOW if active else T.SIDEBAR_HOVER,
                text_color=T.NAVY if active else T.SIDEBAR_TEXT_2,
                image=T.icon(btn._icon_name, 20, T.NAVY if active else T.SIDEBAR_TEXT_2),
            )

        if name == "input":
            self.update_preview_and_validate()
        else:
            # Счётчик строк относится только к списку штрихкодов
            self.char_count_label.configure(text="")

    def create_input_page(self, parent):
        parent.grid_columnconfigure(0, weight=1)
        parent.grid_rowconfigure(2, weight=1)

        # ===== Хедер: надпись, заголовок, плашки-сводка, главные действия =====
        header = ctk.CTkFrame(parent, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", padx=36, pady=(28, 0))

        T.overline(header, "Штрихкоды").pack(anchor="w")
        T.title(header, "Список кодов").pack(anchor="w", pady=(0, 0))

        # Сводка плашками: «7 КОДОВ» «CODE 128» «58×40 ММ» «600 DPI»
        chips = ctk.CTkFrame(header, fg_color="transparent")
        chips.pack(anchor="w", pady=(6, 0))
        self.header_chips = [T.chip(chips, "", strong=(i == 0)) for i in range(5)]
        for c in self.header_chips[:4]:
            c.pack(side="left", padx=(0, 6))

        actions = ctk.CTkFrame(header, fg_color="transparent")
        actions.pack(fill="x", pady=(18, 0))

        self.generate_button = T.button(actions, "Создать PDF", self.generate_pdf,
                                        variant="primary", height=50, width=210,
                                        icon_name="play")
        self.generate_button.pack(side="left")
        T.button(actions, "Предпросмотр", self.show_preview_window, variant="blue",
                 height=50, width=200, icon_name="eye").pack(side="left", padx=(10, 0))
        T.icon_button(actions, "settings", self.show_barcode_settings, size=50,
                      tooltip="Настройки штрихкодов: размер, отступы, качество"
                      ).pack(side="left", padx=(10, 0))

        type_box = ctk.CTkFrame(actions, fg_color="transparent")
        type_box.pack(side="right")
        T.field_label(type_box, "Тип кода").pack(anchor="w", padx=(4, 0))
        self.type_menu = T.option_menu(
            type_box, values=list(BARCODE_TYPE_LABELS.values()),
            command=self.on_type_selected, width=260, height=42,
        )
        self.type_menu.set(BARCODE_TYPE_LABELS.get(self.barcode_type.get(), BARCODE_TYPE_LABELS["code128"]))
        self.type_menu.pack(anchor="w")

        # ===== Панель инструментов списка =====
        toolbar = ctk.CTkFrame(parent, fg_color="transparent")
        toolbar.grid(row=1, column=0, sticky="ew", padx=36, pady=(24, 12))

        for text, cmd in [("Загрузить", self.load_from_file),
                          ("Сохранить", self.save_to_file),
                          ("Случайный код", self.show_random_dialog),
                          ("Excel", self.export_to_excel)]:
            T.button(toolbar, text, cmd, height=36).pack(side="left", padx=(0, 8))
        T.button(toolbar, "Очистить", self.clear_codes, variant="danger",
                 height=36).pack(side="left")

        self.search_count_label = T.label(toolbar, "", size=12, weight="semibold", color=T.TEXT_2)
        self.search_count_label.pack(side="right", padx=(10, 0))
        # Без textvariable: иначе CTkEntry не показывает placeholder
        self.search_entry = T.entry(toolbar, placeholder="Поиск по кодам",
                                    width=230, height=36, corner_radius=18)
        self.search_entry.pack(side="right")
        self.search_entry.bind("<KeyRelease>",
                               lambda e: self.search_var.set(self.search_entry.get()))
        self.search_var.trace_add('write', self.filter_codes)

        # ===== Поле ввода кодов =====
        editor = T.card(parent)
        editor.grid(row=2, column=0, sticky="nsew", padx=36, pady=(0, 22))
        editor.grid_rowconfigure(0, weight=1)
        editor.grid_columnconfigure(0, weight=1)

        self.text_codes = ctk.CTkTextbox(
            editor, font=T.font(15, mono=True), fg_color=T.SURFACE,
            text_color=T.TEXT, corner_radius=16, border_width=0, border_spacing=14,
            wrap="none",
        )
        self.text_codes.grid(row=0, column=0, sticky="nsew", padx=4, pady=4)
        self.text_codes._textbox.configure(
            insertbackground=T.TEXT, selectbackground=T.SELECTION_BG,
            selectforeground=T.TEXT, spacing1=3, spacing3=3,
        )
        ContextMenu(self.root, self.text_codes)
        self._bind_paste_shortcut(self.text_codes)
        self.text_codes.bind("<KeyRelease>", lambda e: self.update_preview_and_validate(), add="+")

        # Пустое состояние поверх поля — «инструкция» из трёх шагов
        self.placeholder = self._build_empty_state(editor)
        self.text_codes.bind("<FocusIn>", lambda e: self._update_placeholder(), add="+")
        self.text_codes.bind("<FocusOut>", lambda e: self._update_placeholder(), add="+")

    def _build_empty_state(self, parent):
        box = ctk.CTkFrame(parent, fg_color=T.SURFACE, corner_radius=0)

        ctk.CTkLabel(box, image=T.eco_illustration(150), text=""
                     ).grid(row=0, column=0, rowspan=5, sticky="n", padx=(0, 30))

        T.label(box, "КАК ЭТО РАБОТАЕТ", size=20, weight="black"
                ).grid(row=0, column=1, sticky="w", pady=(0, 12))
        steps = [
            ("Добавьте коды", "Вставьте их сюда по одному на строку или перетащите файл .txt"),
            ("Выберите тип кода", "Справа вверху. Code 128 подходит почти всегда"),
            ("Нажмите «Создать PDF»", "Готовый файл можно сразу отправить на печать"),
        ]
        for i, (head, text) in enumerate(steps, 1):
            row = ctk.CTkFrame(box, fg_color="transparent")
            row.grid(row=i, column=1, sticky="w", pady=6)
            T.step_badge(row, i, size=32).grid(row=0, column=0, rowspan=2, sticky="n", padx=(0, 12))
            T.label(row, head.upper(), size=13, weight="extrabold", height=18
                    ).grid(row=0, column=1, sticky="w")
            T.label(row, text, size=12, color=T.TEXT_2, height=18
                    ).grid(row=1, column=1, sticky="w")

        # Клик по подсказке — сразу в поле ввода
        def focus_editor(_e=None):
            self.text_codes.focus_set()

        def bind_all(w):
            w.bind("<Button-1>", focus_editor, add="+")
            for child in w.winfo_children():
                bind_all(child)
        bind_all(box)
        box.place(x=30, y=26)
        return box

    def _update_placeholder(self):
        has_text = bool(self.text_codes.get("1.0", "end").strip())
        focused = self.root.focus_get() in (self.text_codes, getattr(self.text_codes, "_textbox", None))
        if has_text or focused:
            self.placeholder.place_forget()
        else:
            self.placeholder.place(x=30, y=26)

    def create_status_bar(self, parent):
        # Чёрная панель внизу — как дисплей на автомате
        bar = ctk.CTkFrame(parent, fg_color=T.DISPLAY, corner_radius=0, height=44)
        bar.grid_propagate(False)

        inner = ctk.CTkFrame(bar, fg_color="transparent")
        inner.pack(fill="both", expand=True, padx=24)

        self.status_dot = T.label(inner, "●", size=13, color=T.SUCCESS)
        self.status_dot.pack(side="left")
        self.status_label = T.label(inner, "Готов к работе", size=13, weight="semibold",
                                    color="#FFFFFF")
        self.status_label.pack(side="left", padx=(8, 0))

        self.char_count_label = T.label(inner, "", size=12, color="#9096A3")
        self.char_count_label.pack(side="right")

        return bar

    # ========== ПОИСК ==========
    def filter_codes(self, *args):
        search_text = self.search_var.get().lower()
        if not search_text:
            self.text_codes.tag_remove("search", "1.0", "end")
            self.search_count_label.configure(text="")
            return

        self.text_codes.tag_remove("search", "1.0", "end")
        count = 0

        content = self.text_codes.get("1.0", "end")
        lines = content.split('\n')

        for i, line in enumerate(lines):
            if search_text in line.lower():
                start = f"{i+1}.0"
                end = f"{i+1}.{len(line)}"
                self.text_codes.tag_add("search", start, end)
                count += 1

        self.text_codes.tag_config("search", background=T.HIGHLIGHT_BG, foreground=T.ACCENT)
        self.search_count_label.configure(text=f"Найдено: {count}")

    # ========== ОБНОВЛЕНИЕ ИНФОРМАЦИИ ==========
    def update_all(self):
        self.update_preview_and_validate()
        self.root.update_idletasks()

    def update_header(self, codes_count=None):
        """Мета-строка под заголовком и подпись главной кнопки."""
        if not hasattr(self, "header_chips"):
            return
        if codes_count is None:
            codes_count = len(self.get_codes_list())
        parts = [
            f"{codes_count} {plural(codes_count, 'код', 'кода', 'кодов')}",
            BARCODE_TYPE_SHORT.get(self.barcode_type.get(), self.barcode_type.get()),
            f"{self.label_width.get()}×{self.label_height.get()} мм",
            f"{self.barcode_dpi.get()} dpi",
            "только выделенные",
        ]
        for chip, text in zip(self.header_chips, parts):
            chip.configure(text=text.upper())
        extra = self.header_chips[4]
        if self.selected_only.get():
            extra.pack(side="left", padx=(0, 6))
        else:
            extra.pack_forget()
        self.generate_button.configure(
            text=" СОЗДАТЬ ZIP" if self.batch_mode.get() else " СОЗДАТЬ PDF")

    def update_preview_and_validate(self, event=None):
        if hasattr(self, "placeholder"):
            self._update_placeholder()

        # Пропускаем валидацию когда активна вкладка QR-паспортов
        if self.current_page == "qr":
            return

        codes = self.get_codes_list()
        self.update_header(len(codes))

        self.text_codes.tag_remove("invalid", "1.0", "end")

        for line_no, code in self._code_lines():
            valid, _ = validate_code(code, self.barcode_type.get())
            if not valid:
                self.text_codes.tag_add("invalid", f"{line_no}.0", f"{line_no}.end")

        self.text_codes.tag_config("invalid", background=T.INVALID_BG, foreground=T.ERROR)

        text = self.text_codes.get("1.0", "end")
        chars = len(text.replace('\n', ''))
        lines = len(codes)
        self.char_count_label.configure(
            text=f"{lines} {plural(lines, 'строка', 'строки', 'строк')}  ·  "
                 f"{chars} {plural(chars, 'символ', 'символа', 'символов')}")

        if codes:
            invalid_count = sum(1 for c in codes if not validate_code(c, self.barcode_type.get())[0])
            if invalid_count > 0:
                self.set_status(f"{invalid_count} из {len(codes)} кодов не подходят для "
                                f"{BARCODE_TYPE_SHORT.get(self.barcode_type.get(), '')} — они подсвечены красным",
                                'warning')
            else:
                self.set_status(f"Готово к печати: {len(codes)} "
                                f"{plural(len(codes), 'штрихкод', 'штрихкода', 'штрихкодов')}", 'success')
        else:
            self.set_status("Добавьте коды, чтобы начать", 'info')

    def set_status(self, text, status_type='success'):
        status_colors = {'success': T.SUCCESS, 'warning': T.WARNING,
                         'error': T.ERROR, 'info': T.TEXT_3}
        # Состояние показывает цветная точка — эмодзи в начале сообщения не нужен
        text = re.sub(r'^[^\w\s«"(]+\s*', '', str(text))
        self.status_dot.configure(text_color=status_colors.get(status_type, T.SUCCESS))
        self.status_label.configure(text=text)
        self.status_label.after(100, self.status_label.update_idletasks)

    def update_status_text(self):
        self.update_preview_and_validate()

    @error_handler
    def on_type_selected(self, choice):
        for val, label in BARCODE_TYPE_LABELS.items():
            if label == choice:
                self.barcode_type.set(val)
                break
        self.settings.set("barcode_type", self.barcode_type.get())
        self.generator.load_settings()
        self.update_all()
        logger.info(f"Выбран тип штрихкода: {self.barcode_type.get()}")

    def sync_type_menu(self):
        """Обновляет выпадающий список типа после изменения в окне настроек."""
        if hasattr(self, "type_menu"):
            self.type_menu.set(BARCODE_TYPE_LABELS.get(self.barcode_type.get(), ""))

    def load_history(self):
        self.history_data = self.settings.get_history()

    # ===== РАБОТА С ФАЙЛАМИ =====
    def _code_lines(self):
        """[(номер строки в поле, код), ...] — пустые строки пропускаются,
        но номера остаются настоящими (нужно для подсветки и выделения)."""
        lines = self.text_codes.get("1.0", "end-1c").split('\n')
        return [(n, line.strip()) for n, line in enumerate(lines, 1) if line.strip()]

    def get_codes_list(self):
        code_lines = self._code_lines()

        if self.selected_only.get():
            selected_ranges = self.text_codes.tag_ranges("sel")
            if selected_ranges:
                selected = set()
                for i in range(0, len(selected_ranges), 2):
                    start_line, _ = map(int, str(selected_ranges[i]).split('.'))
                    end_line, end_col = map(int, str(selected_ranges[i + 1]).split('.'))
                    # Выделение до начала следующей строки её не захватывает
                    if end_col == 0 and end_line > start_line:
                        end_line -= 1
                    selected.update(range(start_line, end_line + 1))
                return [code for n, code in code_lines if n in selected]
        return [code for _, code in code_lines]

    @error_handler
    def load_codes_from_file(self, filename):
        try:
            name = os.path.basename(filename)
            if filename.lower().endswith(('.csv', '.xlsx')):
                # Код + колонки (name, price, ...) для плейсхолдеров шаблона
                rows = read_codes_with_metadata(filename)
                if not rows:
                    messagebox.showwarning("Внимание", f"В файле {name} не найдено кодов")
                    return
                self._append_codes("\n".join(code for code, _ in rows))
                columns = set()
                for code, meta in rows:
                    if meta:
                        self.codes_metadata[code] = meta
                        columns.update(meta)
                self.settings.set("codes_metadata", self.codes_metadata)
                extra = ""
                if columns:
                    extra = " · для шаблона: " + ", ".join("{" + c + "}" for c in sorted(columns))
                message = (f"Добавлено {len(rows)} {plural(len(rows), 'код', 'кода', 'кодов')} "
                           f"из {name}{extra}")
            else:
                with open(filename, 'r', encoding='utf-8') as f:
                    self._append_codes(f.read())
                message = f"Добавлено из {name}"
            # Сначала пересчёт (он пишет свой статус), потом — итог загрузки
            self.update_all()
            self.set_status(message, 'success')
            self.settings.add_recent_file(filename)
            logger.info(f"Загружен файл: {filename}")
        except Exception as e:
            logger.error(f"Ошибка загрузки файла {filename}: {e}")
            raise

    def _append_codes(self, text):
        """Дописывает коды в конец поля, с новой строки."""
        text = text.strip("\n")
        if not text:
            return
        if self.text_codes.get("1.0", "end-1c").strip():
            self.text_codes.insert("end-1c", "\n" + text)
        else:
            self.text_codes.delete("1.0", "end")
            self.text_codes.insert("1.0", text)

    @error_handler
    def load_from_file(self):
        self.show_page("input")
        filename = filedialog.askopenfilename(filetypes=[
            ("Коды и таблицы", "*.txt *.csv *.xlsx"), ("Текстовые файлы", "*.txt"),
            ("CSV", "*.csv"), ("Excel", "*.xlsx")])
        if filename:
            self.load_codes_from_file(filename)

    @error_handler
    def save_to_file(self):
        codes = self.get_codes_list()
        if not codes:
            messagebox.showwarning("Внимание", "Нет кодов для сохранения")
            return
        filename = filedialog.asksaveasfilename(defaultextension=".txt")
        if filename:
            with open(filename, 'w', encoding='utf-8') as f:
                f.write('\n'.join(codes))
            self.set_status("Сохранено", 'success')
            self.settings.add_recent_file(filename)
            logger.info(f"Сохранён файл: {filename}")

    @error_handler
    def clear_codes(self):
        if not self.text_codes.get("1.0", "end").strip():
            return
        if messagebox.askyesno("Подтверждение", "Очистить все коды?"):
            self.text_codes.delete("1.0", "end")
            self.codes_metadata = {}
            self.settings.set("codes_metadata", {})
            self.update_all()
            logger.info("Коды очищены")

    # ===== ДИАЛОГИ =====
    def _dialog(self, title, width, height, modal=True):
        """Пустое окно-диалог в стиле приложения, по центру главного."""
        return T.setup_dialog(ctk.CTkToplevel(self.root), self.root, title, width, height, modal)

    @error_handler
    def show_random_dialog(self):
        """Диалог выбора формата случайного кода."""
        dialog = self._dialog("Случайный код", 440, 450)

        T.title(dialog, "Случайный код", size=22).pack(anchor="w", padx=28, pady=(26, 2))
        T.label(dialog, "Выберите формат — код добавится в конец списка",
                size=13, color=T.TEXT_2).pack(anchor="w", padx=28, pady=(0, 16))

        def _do_generate(make_code):
            self._add_code(make_code())
            dialog.destroy()

        options = [
            ("Серийный номер", "SERIAL-1234", lambda: f"SERIAL-{random.randint(1000, 9999)}"),
            ("Короткий UUID", "a1b2c3d4", lambda: str(uuid.uuid4())[:8]),
            ("Полный UUID", "a1b2c3d4-e5f6-…", lambda: str(uuid.uuid4())),
        ]

        for name, example, make in options:
            row = T.card(dialog, corner_radius=14, cursor="hand2")
            row.pack(fill="x", padx=24, pady=4)
            T.label(row, name, size=14, weight="bold").pack(anchor="w", padx=18, pady=(12, 0))
            ctk.CTkLabel(row, text=example, text_color=T.TEXT_2, font=T.font(12, mono=True)
                         ).pack(anchor="w", padx=18, pady=(0, 12))
            T.bind_hover(row, T.SURFACE, T.SURFACE_2)
            for w in [row] + row.winfo_children():
                w.bind("<Button-1>", lambda e, m=make: _do_generate(m))

        T.button(dialog, "Отмена", dialog.destroy, variant="ghost", width=120
                 ).pack(pady=(14, 20))

    @error_handler
    def _add_code(self, code):
        """Добавляет код в текстовое поле и обновляет превью."""
        self.show_page("input")
        current = self.text_codes.get("1.0", "end").strip()
        new_text = current + "\n" + code if current else code
        self.text_codes.delete("1.0", "end")
        self.text_codes.insert("1.0", new_text)
        self.update_all()
        self.set_status(f"Добавлен код: {code}", 'success')
        logger.info(f"Сгенерирован случайный код: {code}")

    @error_handler
    def generate_random_code(self):
        """Старое имя метода — оставлено для совместимости с горячей клавишей Ctrl+R."""
        self.show_random_dialog()

    @error_handler
    def show_template_dialog(self):
        TemplateEditor(self.root, self)
        logger.info("Открыт редактор шаблона этикетки")

    @error_handler
    def export_to_excel(self):
        try:
            from openpyxl import Workbook
            import openpyxl.styles
        except ImportError:
            messagebox.showwarning("Внимание", "Установите openpyxl: pip install openpyxl")
            return

        codes = self.get_codes_list()
        if not codes:
            messagebox.showwarning("Внимание", "Нет кодов для экспорта")
            return

        filename = filedialog.asksaveasfilename(title="Сохранить Excel", defaultextension=".xlsx",
                                                  filetypes=[("Excel файлы", "*.xlsx")])
        if not filename:
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Штрихкоды"

            headers = ["№", "Код", "Тип", "Статус", "Дата генерации"]
            for col, header in enumerate(headers, 1):
                ws.cell(row=1, column=col, value=header)
                ws.cell(row=1, column=col).font = openpyxl.styles.Font(bold=True)

            for i, code in enumerate(codes, 1):
                ws.cell(row=i + 1, column=1, value=i)
                ws.cell(row=i + 1, column=2, value=code)
                ws.cell(row=i + 1, column=3, value=self.barcode_type.get())
                ws.cell(row=i + 1, column=4, value="✓ OK")
                ws.cell(row=i + 1, column=5, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

            wb.save(filename)
            self.set_status(f"Экспортировано в Excel: {os.path.basename(filename)}", 'success')
            messagebox.showinfo("Готово", f"Экспорт завершён!\n{filename}")
            logger.info(f"Экспорт в Excel: {filename}")
        except Exception as e:
            logger.error(f"Ошибка экспорта в Excel: {e}")
            raise

    @error_handler
    def show_history_window(self):
        self.load_history()
        window = self._dialog("История", 760, 560, modal=False)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(1, weight=1)

        head = ctk.CTkFrame(window, fg_color="transparent")
        head.grid(row=0, column=0, sticky="ew", padx=28, pady=(26, 14))
        T.overline(head, "Последние 50 генераций").pack(anchor="w")
        T.title(head, "История", size=26).pack(anchor="w")

        body = ctk.CTkScrollableFrame(window, fg_color="transparent")
        body.grid(row=1, column=0, sticky="nsew", padx=16, pady=(0, 8))

        if not self.history_data:
            T.label(body, "Здесь появятся созданные PDF", size=14, color=T.TEXT_3
                    ).pack(pady=60)
        else:
            for i, item in enumerate(reversed(self.history_data), 1):
                row = ctk.CTkFrame(body, fg_color="transparent", corner_radius=12)
                row.pack(fill="x", pady=1)
                row.grid_columnconfigure(1, weight=1)

                T.label(row, str(i), size=13, color=T.TEXT_3, width=28
                        ).grid(row=0, column=0, rowspan=2, padx=(10, 8), pady=8)
                T.label(row, item.get('file', ''), size=14, weight="bold", anchor="w"
                        ).grid(row=0, column=1, sticky="w", pady=(8, 0))
                count = item.get('codes_count', 0)
                T.label(row, f"{count} {plural(count, 'этикетка', 'этикетки', 'этикеток')}  ·  "
                             f"{item.get('barcode_type', '')}  ·  {item.get('label_size', '')} мм",
                        size=12, color=T.TEXT_2, anchor="w"
                        ).grid(row=1, column=1, sticky="w", pady=(0, 8))
                T.label(row, item.get('time', ''), size=12, color=T.TEXT_3
                        ).grid(row=0, column=2, rowspan=2, padx=14)
                T.bind_hover(row, "transparent", T.SURFACE)

        T.button(window, "Закрыть", window.destroy, width=130
                 ).grid(row=2, column=0, sticky="e", padx=28, pady=(4, 20))
        logger.info("Открыто окно истории")

    @error_handler
    def show_preview_window(self):
        codes = self.get_codes_list()
        if not codes:
            messagebox.showwarning("Внимание", "Нет кодов для предпросмотра")
            return

        try:
            page_w = float(self.label_width.get())
            page_h = float(self.label_height.get())
            if page_w <= 0 or page_h <= 0:
                raise ValueError("Размер должен быть больше 0")
        except ValueError as e:
            messagebox.showerror("Ошибка", f"Введите корректные числовые значения\n{str(e)}")
            return

        self.save_current_settings()
        self.generator.load_settings()
        self.set_status("Готовим предпросмотр...", 'warning')
        PreviewWindow(self.root, codes, self)  # статус «готов» выставит само окно
        logger.info("Открыто окно предпросмотра печати")

    # ===== ГЕНЕРАЦИЯ PDF =====
    @error_handler
    def generate_pdf(self):
        if self._job is not None:
            self.set_status("Генерация уже идёт", 'warning')
            return

        codes = self.get_codes_list()
        if not codes:
            messagebox.showwarning("Внимание", "Нет кодов для генерации")
            return

        try:
            page_w = float(self.label_width.get())
            page_h = float(self.label_height.get())
            if page_w <= 0 or page_h <= 0:
                raise ValueError("Размер должен быть больше 0")
        except ValueError as e:
            messagebox.showerror("Ошибка", f"Введите корректные числовые значения\n{str(e)}")
            return

        self.save_current_settings()
        self.generator.load_settings()

        as_zip = self.batch_mode.get()
        if as_zip:
            filename = filedialog.asksaveasfilename(title="Сохранить ZIP архив", defaultextension=".zip",
                                                    filetypes=[("ZIP архивы", "*.zip")], initialfile="barcodes.zip")
        else:
            filename = filedialog.asksaveasfilename(title="Сохранить PDF", defaultextension=".pdf",
                                                    filetypes=[("PDF файлы", "*.pdf")], initialfile="barcodes.pdf")
        if not filename:
            return

        job = {
            "queue": queue.Queue(),
            "cancel": threading.Event(),
            "filename": filename,
            "as_zip": as_zip,
            "total": len(codes),
            "barcode_type": self.barcode_type.get(),
            "label_size": f"{self.label_width.get()}×{self.label_height.get()}",
        }
        self._job = job
        self.show_progress(len(codes), "Создание ZIP" if as_zip else "Создание PDF")
        self.set_status("Генерация штрихкодов...", 'warning')
        threading.Thread(target=self._generate_worker, args=(job, codes, dict(self.codes_metadata)),
                         daemon=True).start()
        self.root.after(50, self._poll_job)

    def _generate_worker(self, job, codes, metadata):
        """Фоновый поток генерации. К Tk не обращается — только кладёт
        сообщения в очередь, которую забирает _poll_job в основном потоке."""
        q = job["queue"]
        try:
            with tempfile.TemporaryDirectory() as temp_dir:
                images, invalid = [], []
                for i, code in enumerate(codes, 1):
                    if job["cancel"].is_set():
                        q.put(("cancelled", None))
                        return
                    valid, msg = validate_code(code, job["barcode_type"])
                    if not valid:
                        invalid.append(f"{code} — {msg}")
                    else:
                        img_path = self.generator.generate_barcode_image(code, temp_dir)
                        if img_path:
                            images.append((code, img_path))
                        else:
                            invalid.append(f"{code} — ошибка генерации")
                    q.put(("progress", i))

                if job["cancel"].is_set():
                    q.put(("cancelled", None))
                    return
                if not images:
                    q.put(("empty", invalid))
                    return

                q.put(("saving", None))
                if job["as_zip"]:
                    self.generator.write_zip(job["filename"], images)
                else:
                    self.generator.create_pdf(job["filename"], images, metadata=metadata)
                q.put(("done", (len(images), invalid)))
        except Exception as e:
            logger.error(f"Ошибка генерации:\n{traceback.format_exc()}")
            q.put(("error", e))

    def _poll_job(self):
        """Забирает сообщения фоновой генерации и обновляет окно прогресса."""
        job = self._job
        if job is None:
            return
        try:
            while True:
                kind, data = job["queue"].get_nowait()
                if kind == "progress":
                    self.update_progress(data, job["total"])
                elif kind == "saving":
                    self.progress_label.configure(text="Сохранение файла…")
                else:
                    self._finish_job(job, kind, data)
                    return
        except queue.Empty:
            pass
        self.root.after(50, self._poll_job)

    def _finish_job(self, job, kind, data):
        self._job = None
        self.hide_progress()
        name = os.path.basename(job["filename"])

        if kind == "cancelled":
            self.set_status("Генерация отменена", 'warning')
            return
        if kind == "error":
            self.set_status("Ошибка генерации", 'error')
            messagebox.showerror("Ошибка", f"Не удалось создать файл:\n{data}")
            return

        invalid = data if kind == "empty" else data[1]
        if invalid:
            msg = "Следующие коды не были сгенерированы:\n" + "\n".join(invalid[:10])
            if len(invalid) > 10:
                msg += f"\n... и ещё {len(invalid) - 10}"
            logger.warning(f"Не сгенерированы коды: {invalid}")
        if kind == "empty":
            self.set_status("Ошибка генерации", 'error')
            messagebox.showerror("Ошибка", "Не удалось сгенерировать ни одного штрихкода"
                                 + ("\n\n" + msg if invalid else ""))
            return
        if invalid:
            messagebox.showwarning("Предупреждение", msg)

        count = data[0]
        if job["as_zip"]:
            self.set_status(f"ZIP архив создан: {name}", 'success')
            messagebox.showinfo("Готово", f"ZIP архив создан!\nСохранено {count} изображений")
            logger.info(f"Создан ZIP архив: {job['filename']} ({count} изображений)")
            return

        self.settings.add_history(job["filename"], count, job["barcode_type"], job["label_size"])
        self.load_history()
        self.set_status(f"PDF создан: {name}", 'success')
        logger.info(f"Создан PDF: {job['filename']} ({count} этикеток)")
        if messagebox.askyesno("Готово", f"PDF создан!\n{job['filename']}\n\nОткрыть файл?"):
            open_pdf(job["filename"])

    def cancel_generation(self):
        if self._job is not None:
            self._job["cancel"].set()
            if self.progress_window:
                self.progress_label.configure(text="Отмена…")

    def save_current_settings(self):
        self.settings.update({
            "barcode_type": self.barcode_type.get(),
            "barcode_label_width": self.label_width.get(),
            "barcode_label_height": self.label_height.get(),
            "barcode_margin_top": self.margin_top.get(),
            "barcode_margin_bottom": self.margin_bottom.get(),
            "barcode_margin_left": self.margin_left.get(),
            "barcode_margin_right": self.margin_right.get(),
            "barcode_show_border": self.show_border.get(),
            "barcode_labels_per_page": self.labels_per_page.get(),
            "barcode_selected_only": self.selected_only.get(),
            "barcode_batch_mode": self.batch_mode.get(),
            "barcode_use_template": self.use_template.get(),
            "barcode_show_text": self.show_text.get(),
            "barcode_dpi": self.barcode_dpi.get(),
        })

    @error_handler
    def show_barcode_settings(self):
        """Открывает окно настроек штрихкодов."""
        from .gui_barcode_settings import BarcodeSettingsWindow
        from .gui_widgets import fix_transparency_after_close

        BarcodeSettingsWindow(self.root, self)
        # Когда окно закроется — принудительно убираем прозрачность главного
        self.root.after(100, lambda: fix_transparency_after_close(self.root))

    def show_progress(self, total, title_text="Генерация"):
        self.progress_window = self._dialog(title_text, 440, 176)
        self.progress_window.resizable(False, False)
        self.progress_window.protocol("WM_DELETE_WINDOW", self.cancel_generation)

        T.label(self.progress_window, title_text.upper(), size=17, weight="black"
                ).pack(anchor="w", padx=28, pady=(26, 14))

        self.progress_bar = ctk.CTkProgressBar(self.progress_window, height=10, corner_radius=5)
        self.progress_bar.set(0)
        self.progress_bar.pack(fill="x", padx=28)

        row = ctk.CTkFrame(self.progress_window, fg_color="transparent")
        row.pack(fill="x", padx=28, pady=(14, 0))
        self.progress_label = T.label(row, f"0 из {total}", size=13, weight="semibold", color=T.TEXT_2)
        self.progress_label.pack(side="left")
        T.button(row, "Отмена", self.cancel_generation, variant="danger", height=36, width=110
                 ).pack(side="right")

    def update_progress(self, current, total):
        if self.progress_bar:
            self.progress_bar.set(current / max(total, 1))
            self.progress_label.configure(text=f"{current} из {total}")

    def hide_progress(self):
        if self.progress_window:
            try:
                self.progress_window.grab_release()
            except Exception:
                pass
            self.progress_window.destroy()
            self.progress_window = None
            self.progress_bar = None
